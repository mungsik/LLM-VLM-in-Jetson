"""Phi-4-mini GGUF 양자화 스윕(8종)을 엑셀로: 요약(비트포함)/과목별/문항별비교."""
import json, os
from collections import defaultdict, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = "docs/results/quant_sweep.jsonl"
KO_IFEVAL_SRC = "docs/results/ko_ifeval.jsonl"
OUT = "docs/results/2026-06-21-phi4mini-양자화스윕.xlsx"

# quant -> (실효 bpw, 파일크기 GB)  ※ HfApi files_metadata 측정값, params=3.836B 기준
META = OrderedDict([
    ("BF16",   (16.02, 7.68)),
    ("Q8_0",   ( 8.52, 4.08)),
    ("Q6_K",   ( 6.58, 3.16)),
    ("Q5_K_M", ( 5.94, 2.85)),
    ("Q4_K_M", ( 5.20, 2.49)),
    ("Q3_K_M", ( 4.42, 2.12)),
    ("Q2_K_L", ( 3.51, 1.68)),
    ("Q2_K",   ( 3.51, 1.68)),
])
# 한국어 PPL (llama.cpp llama-perplexity, kowiki, ctx=512). 낮을수록 좋음.
# ※ transformers gguf 로더는 Phi-4-mini LongRoPE 오변환으로 부정확 → llama.cpp 사용.
PPL = {
    "BF16": 12.19, "Q8_0": 12.20, "Q6_K": 12.40, "Q5_K_M": 12.51,
    "Q4_K_M": 12.75, "Q3_K_M": 14.71, "Q2_K_L": 63.61, "Q2_K": 63.61,
}
ORDER = list(META.keys())  # bpw 큰 순

rows = [json.loads(l) for l in open(SRC, encoding="utf-8")]
def quant_of(r): return r["model"].replace("phi4mini-", "")
QUANTS = [q for q in ORDER if any(quant_of(r) == q for r in rows)]

def metric(metrics, name):
    if name in metrics:
        return metrics[name]
    for k, v in metrics.items():
        if k.split(",")[0] == name:
            return v
    return None

def load_ifeval(path):
    out = {}
    if not os.path.exists(path):
        return out
    with open(path, encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue
            r = json.loads(line)
            q = r["name"].replace("phi4mini-", "")
            m = r.get("metrics", {})
            out[q] = {
                "prompt_strict": metric(m, "prompt_level_strict_acc"),
                "prompt_loose": metric(m, "prompt_level_loose_acc"),
                "inst_strict": metric(m, "inst_level_strict_acc"),
                "inst_loose": metric(m, "inst_level_loose_acc"),
            }
    return out

def pct(v):
    return f"{v*100:.1f}%" if isinstance(v, (int, float)) else "-"

KO_IFEVAL = load_ifeval(KO_IFEVAL_SRC)

HEAD_FILL = PatternFill("solid", fgColor="305496"); HEAD_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="C6EFCE"); NO_FILL = PatternFill("solid", fgColor="FFC7CE")
TITLE = Font(bold=True, size=13)
CENTER = Alignment(horizontal="center", vertical="center"); WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D0D0D0"); BORDER = Border(THIN, THIN, THIN, THIN)

def header(ws, names, row=1):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n); c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.alignment = CENTER; c.border = BORDER
def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[get_column_letter(i)].width = w

# 정확도 집계
agg = {q: [0, 0] for q in QUANTS}
for r in rows:
    q = quant_of(r)
    if q in agg: agg[q][1] += 1; agg[q][0] += 1 if r["correct"] else 0
acc = {q: (agg[q][0] / agg[q][1] * 100 if agg[q][1] else 0) for q in QUANTS}
ref = acc.get("BF16", 0)

wb = Workbook()

# ---- Sheet 1: 양자화 스윕 요약 ----
ws = wb.active; ws.title = "양자화스윕"
ws["A1"] = "Phi-4-mini-instruct GGUF 양자화 스윕 — KMMLU + PPL + Ko-IFEval"
ws["A1"].font = TITLE; ws.merge_cells("A1:K1")
header(ws, ["quant", "실효 비트(bpw)", "파일크기(GB)", "KMMLU", "정답/900",
           "PPL(한국어)", "PPL vs BF16", "Ko-IFEval strict", "Ko-IFEval loose",
           "Jetson 8GB", "권장"], row=3)
ppl_ref = PPL.get("BF16", 0)
rr = 4
for q in QUANTS:
    bpw, gb = META[q]
    p = PPL.get(q)
    ppl_delta = f"+{(p/ppl_ref-1)*100:.0f}%" if p and ppl_ref else "-"
    fit = "여유" if gb < 6 else ("빠듯" if gb < 8 else "불가")
    # 권장: PPL이 BF16 대비 +10% 이내면 안전
    rec = "✓ 안전" if (p and p <= ppl_ref*1.10) else ("△ 주의" if (p and p <= ppl_ref*1.5) else "✗ 붕괴")
    ko = KO_IFEVAL.get(q, {})
    vals = [q, bpw, gb, f"{acc[q]:.2f}%", f"{agg[q][0]}/900",
            f"{p:.2f}" if p else "-", ppl_delta,
            pct(ko.get("prompt_strict")), pct(ko.get("prompt_loose")), fit, rec]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=rr, column=c, value=v); cell.alignment = CENTER; cell.border = BORDER
    rr += 1
widths(ws, [10, 13, 12, 9, 9, 11, 11, 15, 15, 10, 9])
ws.append([])
for note in [
    "※ 실효 bpw = 파일비트/파라미터수. 명목(Q4=4bit)보다 높은 건 임베딩·출력층을 고정밀로 남기기 때문.",
    "※ Phi-4-mini = 3.836B. BF16(비양자화)이 기준점.",
    "※ KMMLU(객관식 logprob)는 양자화 손실을 거의 못 잡음(바닥점수+argmax robust) — quant 간 평평.",
    "※ PPL(llama.cpp, 한국어 wiki, ctx512)이 진짜 손실을 드러냄: Q4까지 무손실, Q3부터 열화, Q2 붕괴.",
    "※ PPL은 반드시 llama.cpp로 측정 — transformers gguf 로더는 Phi-4-mini LongRoPE 오변환으로 부정확.",
    "※ Ko-IFEval은 docs/results/ko_ifeval.jsonl 이 있을 때 자동 병합. 현재 미측정 모델은 '-'로 표시.",
    "※ Q2_K / Q2_K_L 파일크기·PPL 동일하게 측정됨.",
]:
    ws.append(["", note])

# ---- Sheet 2: Ko-IFEval ----
ws_if = wb.create_sheet("Ko-IFEval")
ws_if["A1"] = "Ko-IFEval / IFEval 생성형 지시 준수 평가"; ws_if["A1"].font = TITLE
ws_if.merge_cells("A1:F1")
header(ws_if, ["quant", "prompt strict", "prompt loose", "inst strict", "inst loose", "상태"], row=3)
rr = 4
for q in QUANTS:
    ko = KO_IFEVAL.get(q, {})
    vals = [
        q,
        pct(ko.get("prompt_strict")),
        pct(ko.get("prompt_loose")),
        pct(ko.get("inst_strict")),
        pct(ko.get("inst_loose")),
        "측정됨" if ko else "미측정",
    ]
    for c, v in enumerate(vals, 1):
        cell = ws_if.cell(row=rr, column=c, value=v); cell.alignment = CENTER; cell.border = BORDER
    rr += 1
widths(ws_if, [10, 13, 13, 13, 13, 12])
ws_if.append([])
ws_if.append(["", "입력 파일", KO_IFEVAL_SRC if KO_IFEVAL else f"{KO_IFEVAL_SRC} 없음"])
ws_if.append(["", "권장 실행", "compression/benchmarks/run_ko_ifeval_sweep.sh"])

# ---- Sheet 3: 과목별 ----
ws2 = wb.create_sheet("과목별")
ws2["A1"] = "과목별 정확도 (quant 비교)"; ws2["A1"].font = TITLE
ws2.merge_cells(f"A1:{get_column_letter(len(QUANTS)+1)}1")
header(ws2, ["과목"] + QUANTS, row=3)
subj = defaultdict(lambda: {q: [0, 0] for q in QUANTS})
for r in rows:
    q = quant_of(r)
    if q in QUANTS: s = subj[r["subject"]][q]; s[1] += 1; s[0] += 1 if r["correct"] else 0
rr = 4
for sub in sorted(subj):
    ws2.cell(row=rr, column=1, value=sub).border = BORDER
    for j, q in enumerate(QUANTS, 2):
        ok, tot = subj[sub][q]
        cell = ws2.cell(row=rr, column=j, value=f"{ok/tot*100:.0f}%" if tot else "-")
        cell.alignment = CENTER; cell.border = BORDER
    rr += 1
widths(ws2, [30] + [9]*len(QUANTS))

# ---- Sheet 4: 문항별 비교 ----
ws3 = wb.create_sheet("문항별비교")
ws3["A1"] = "문항별 예측 비교 — quant별 답 (●정답 ✗오답)"; ws3["A1"].font = TITLE
ws3.merge_cells(f"A1:{get_column_letter(7+len(QUANTS))}1")
header(ws3, ["과목", "질문", "A", "B", "C", "D", "정답"] + QUANTS, row=3)
piv = OrderedDict()
for r in rows:
    q = quant_of(r)
    if q not in QUANTS: continue
    k = (r["subject"], r["question"])
    if k not in piv:
        piv[k] = {"A": r["A"], "B": r["B"], "C": r["C"], "D": r["D"], "gold": r["gold"], "p": {}}
    piv[k]["p"][q] = (r["pred"], r["correct"])
rr = 4
for (sub, qn), v in piv.items():
    for c, val in enumerate([sub, qn, v["A"], v["B"], v["C"], v["D"], v["gold"]], 1):
        ws3.cell(row=rr, column=c, value=val)
    for j, q in enumerate(QUANTS, 8):
        pred, ok = v["p"].get(q, ("-", None))
        cell = ws3.cell(row=rr, column=j, value=(f"{pred} {'●' if ok else '✗'}" if ok is not None else "-"))
        if ok is not None: cell.fill = OK_FILL if ok else NO_FILL
        cell.alignment = CENTER
    for c in range(1, 8+len(QUANTS)):
        ws3.cell(row=rr, column=c).border = BORDER
        ws3.cell(row=rr, column=c).alignment = WRAP if c == 2 else CENTER
    rr += 1
widths(ws3, [20, 52, 11, 11, 11, 11, 6] + [9]*len(QUANTS))
ws3.freeze_panes = "A4"

wb.save(OUT)
print("saved:", OUT, "| quants:", QUANTS)
for q in QUANTS:
    print(f"  {q:8s} {META[q][0]:5.2f}bpw  {acc[q]:.2f}%")

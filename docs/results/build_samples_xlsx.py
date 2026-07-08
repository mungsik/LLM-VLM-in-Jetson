"""samples.jsonl(+gguf) 문항별 예측을 엑셀로 변환: 요약/과목별/문항비교/모델별."""
import json, os
from collections import defaultdict, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRCS = ["docs/results/samples.jsonl", "docs/results/gguf_samples.jsonl"]
OUT = "docs/results/2026-06-17-kmmlu-문항별예측.xlsx"

# 표시 순서 + 한글 라벨 (체급 표기 포함)
MODEL_KR = OrderedDict([
    ("original", "원본Phi-4(14.7B)"),
    ("depth", "depth(~10.5B)"),
    ("act", "활성값width(~11B)"),
    ("mag", "magnitude(~11B)"),
    ("phi4mini-q4km", "Phi4mini-Q4(3.8B)"),
])

rows = []
for s in SRCS:
    if os.path.exists(s):
        rows += [json.loads(l) for l in open(s, encoding="utf-8")]

MODELS = [m for m in MODEL_KR if any(r["model"] == m for r in rows)]

HEAD_FILL = PatternFill("solid", fgColor="305496")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
NO_FILL = PatternFill("solid", fgColor="FFC7CE")
TITLE = Font(bold=True, size=13)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, names, row=1):
    for i, name in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.fill = HEAD_FILL; c.font = HEAD_FONT; c.alignment = CENTER; c.border = BORDER


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------- Sheet 1: 요약 ----------
ws = wb.active; ws.title = "요약"
ws["A1"] = "KMMLU 문항별 예측 평가 요약 (limit=20, 과목당 20문항 × 45과목 = 900/모델)"
ws["A1"].font = TITLE; ws.merge_cells("A1:E1")
header(ws, ["모델", "설명", "정답 수", "총 문항", "정확도"], row=3)
agg = OrderedDict((m, [0, 0]) for m in MODELS)
for r in rows:
    if r["model"] in agg:
        agg[r["model"]][1] += 1
        agg[r["model"]][0] += 1 if r["correct"] else 0
rr = 4
for m in MODELS:
    ok, tot = agg[m]
    for c, v in enumerate([m, MODEL_KR[m], ok, tot, f"{ok/tot*100:.2f}%" if tot else "-"], 1):
        cell = ws.cell(row=rr, column=c, value=v); cell.border = BORDER; cell.alignment = CENTER
    rr += 1
widths(ws, [16, 20, 10, 10, 10])

# ---------- Sheet 2: 과목별 정확도 ----------
ws2 = wb.create_sheet("과목별정확도")
ws2["A1"] = "과목별 정확도 (모델 비교)"; ws2["A1"].font = TITLE
ws2.merge_cells(f"A1:{get_column_letter(len(MODELS)+2)}1")
header(ws2, ["과목"] + [MODEL_KR[m] for m in MODELS] + ["문항수"], row=3)
subj = defaultdict(lambda: {m: [0, 0] for m in MODELS})
for r in rows:
    if r["model"] in MODELS:
        s = subj[r["subject"]][r["model"]]; s[1] += 1; s[0] += 1 if r["correct"] else 0
rr = 4
for sub in sorted(subj):
    ws2.cell(row=rr, column=1, value=sub).border = BORDER
    n = 0
    for j, m in enumerate(MODELS, 2):
        ok, tot = subj[sub][m]; n = max(n, tot)
        cell = ws2.cell(row=rr, column=j, value=f"{ok/tot*100:.0f}%" if tot else "-")
        cell.alignment = CENTER; cell.border = BORDER
    c = ws2.cell(row=rr, column=len(MODELS)+2, value=n); c.alignment = CENTER; c.border = BORDER
    rr += 1
widths(ws2, [30] + [18]*len(MODELS) + [9])

# ---------- Sheet 3: 문항별 비교 ----------
ws3 = wb.create_sheet("문항별비교")
ws3["A1"] = "문항별 예측 비교 — 같은 문항에 대한 각 모델의 답 (●정답 ✗오답)"
ws3["A1"].font = TITLE; ws3.merge_cells(f"A1:{get_column_letter(7+len(MODELS))}1")
cols = ["과목", "질문", "A", "B", "C", "D", "정답"] + [MODEL_KR[m]+"답" for m in MODELS]
header(ws3, cols, row=3)
piv = OrderedDict()
for r in rows:
    if r["model"] not in MODELS:
        continue
    key = (r["subject"], r["question"])
    if key not in piv:
        piv[key] = {"A": r["A"], "B": r["B"], "C": r["C"], "D": r["D"],
                    "gold": r["gold"], "preds": {}}
    piv[key]["preds"][r["model"]] = (r["pred"], r["correct"])
rr = 4
for (sub, q), v in piv.items():
    for c, val in enumerate([sub, q, v["A"], v["B"], v["C"], v["D"], v["gold"]], 1):
        ws3.cell(row=rr, column=c, value=val)
    for j, m in enumerate(MODELS, 8):
        pred, ok = v["preds"].get(m, ("-", None))
        cell = ws3.cell(row=rr, column=j, value=(f"{pred} {'●' if ok else '✗'}" if ok is not None else "-"))
        if ok is not None:
            cell.fill = OK_FILL if ok else NO_FILL
        cell.alignment = CENTER
    for c in range(1, 8+len(MODELS)):
        ws3.cell(row=rr, column=c).border = BORDER
        ws3.cell(row=rr, column=c).alignment = WRAP if c == 2 else (CENTER if c >= 3 else WRAP)
    rr += 1
widths(ws3, [22, 58, 13, 13, 13, 13, 7] + [16]*len(MODELS))
ws3.freeze_panes = "A4"

# ---------- Sheets 4+: 모델별 원본 ----------
for m in MODELS:
    safe = m.replace("/", "_")[:25]
    wsm = wb.create_sheet(f"raw-{safe}")
    header(wsm, ["과목", "질문", "A", "B", "C", "D", "정답", "모델답", "정오"], row=1)
    rr = 2
    for r in rows:
        if r["model"] != m:
            continue
        wsm.cell(row=rr, column=1, value=r["subject"])
        wsm.cell(row=rr, column=2, value=r["question"]).alignment = WRAP
        for j, k in enumerate(["A", "B", "C", "D"], 3):
            wsm.cell(row=rr, column=j, value=r[k]).alignment = CENTER
        wsm.cell(row=rr, column=7, value=r["gold"]).alignment = CENTER
        wsm.cell(row=rr, column=8, value=r["pred"]).alignment = CENTER
        oc = wsm.cell(row=rr, column=9, value="O" if r["correct"] else "X"); oc.alignment = CENTER
        oc.fill = OK_FILL if r["correct"] else NO_FILL
        rr += 1
    widths(wsm, [22, 60, 14, 14, 14, 14, 7, 9, 7])
    wsm.freeze_panes = "A2"

wb.save(OUT)
print("saved:", OUT, "| models:", MODELS, "| rows:", len(rows))

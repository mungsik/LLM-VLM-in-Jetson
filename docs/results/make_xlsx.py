"""프루닝 실험 결과 + Q&A 정리를 엑셀(.xlsx)로 출력."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

HEAD_FILL = PatternFill("solid", fgColor="305496")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
TITLE_FONT = Font(bold=True, size=13)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------- Sheet 1: 프루닝 결과 비교 ----------------
ws = wb.active
ws.title = "프루닝 결과"
ws["A1"] = "Phi-4 구조적 프루닝 방법 비교 (KMMLU, limit=20, distillation 회복 전)"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:G1")

header = ["방법(산출물)", "축", "중요도 기준", "감축률", "KMMLU acc", "원본 대비", "구현"]
ws.append([])  # row2 spacer
ws.append(header)
style_header(ws, 3, len(header))

rows = [
    ["원본 Phi-4 (microsoft/phi-4)", "—", "—", "—", "34.11%", "100%", "—"],
    ["depth (phi4-pruned-depth)", "depth", "Block Influence", "27.9%", "31.89%", "~94% (최고)", "직접 구현 (ShortGPT)"],
    ["width (phi4-pruned-act)", "width(MLP)", "활성값 L2", "22.5%", "18.33%", "~54%", "직접 구현 (Minitron)"],
    ["width (phi4-pruned-smoke)", "width(MLP)", "weight |W|", "22.5%", "7.33%", "~21%", "직접 구현 (magnitude)"],
]
for r in rows:
    ws.append(r)
# 최고 행(depth) 강조
for c in range(1, len(header) + 1):
    ws.cell(row=5, column=c).fill = BEST_FILL
for r in range(4, 8):
    for c in range(1, len(header) + 1):
        ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=c).alignment = Alignment(vertical="center")

ws.append([])
note = [
    "메모리(4bit 추정): 원본 7.33GB → depth 5.28GB / width 5.68GB",
    "랜덤 베이스라인(4지선다) ≈ 25%",
    "caveat: limit=20(과목당 20문항, 총 900문항)이라 표본 작아 노이즈 있음 → 방향성은 신뢰, 절대값은 잠정치",
    "전부 distillation 회복 전 수치. depth는 회복 거의 불필요(94%), width는 회복 필요.",
]
for n in note:
    ws.append(["※ " + n])
widths(ws, [30, 12, 16, 10, 12, 14, 22])

# ---------------- Sheet 2: 실험 설정 / 데이터 ----------------
ws2 = wb.create_sheet("실험설정·데이터")
ws2["A1"] = "실험 환경 / 설정 / 데이터"
ws2["A1"].font = TITLE_FONT
ws2.merge_cells("A1:C1")
ws2.append([])
ws2.append(["항목", "값", "비고"])
style_header(ws2, 3, 3)
setup = [
    ["GPU", "RTX PRO 6000 Blackwell 96GB", "GCP phi4-blackwell, g4-standard-48, asia-east1-a"],
    ["드라이버 / torch", "580 / 2.12.0+cu130", "Blackwell sm_120 지원"],
    ["모델", "microsoft/phi-4 (14.66B)", "Phi-3 아키텍처, fused gate_up_proj"],
    ["보정셋(calibration)", "한국어 3종, 256샘플 × 1024토큰", "KoCommercial / KoAlpaca-RealQA / kowikitext-qa"],
    ["감축 목표(width_ratio)", "0.30", "MLP intermediate 30% 감축"],
    ["평가셋", "HAERAE-HUB/KMMLU (test)", "45과목 4지선다, lm-eval 자동 다운로드"],
    ["평가 하네스", "EleutherAI lm-evaluation-harness 0.4.12", "simple_evaluate(model='hf')"],
    ["평가 범위", "limit=20 (과목당 20문항, 총 900문항)", "전체 ~35,000문항의 약 2~3%"],
    ["평가 방식", "0-shot, multiple_choice", "A/B/C/D logprob 최댓값 선택"],
    ["브랜치", "recovery-task0-6", "fork(mungsik/LLM-VLM-in-Jetson)에 push, 단위테스트 15 passed"],
]
for r in setup:
    ws2.append(r)
    for c in range(1, 4):
        ws2.cell(row=ws2.max_row, column=c).border = BORDER
        ws2.cell(row=ws2.max_row, column=c).alignment = WRAP
widths(ws2, [22, 40, 48])

# ---------------- Sheet 3: 파이프라인 (논문 vs 우리) ----------------
ws3 = wb.create_sheet("파이프라인")
ws3["A1"] = "Minitron(2408.11796) 파이프라인 vs 우리 구현 현황"
ws3["A1"].font = TITLE_FONT
ws3.merge_cells("A1:D1")
ws3.append([])
ws3.append(["단계", "출처", "내용", "우리 구현"])
style_header(ws3, 3, 4)
pipe = [
    ["① teacher correction", "논문(2408.11796)", "원본을 distillation 데이터에 파인튜닝 → teacher 생성 (프루닝 前)", "미구현"],
    ["② 중요도 추정", "Minitron", "활성값 기반(width: 뉴런/head/embedding), depth: Block Influence", "MLP 뉴런 + depth(BI)만"],
    ["③ 프루닝", "Minitron", "one-shot 구조적. width(MLP) / depth 둘 다", "구현 완료"],
    ["④ distillation", "Minitron", "teacher(corrected)→student(pruned), logit KD(KL)", "미구현 (PR3)"],
    ["⑤ 양자화", "우리 추가(논문엔 없음)", "GGUF imatrix Q4→IQ3/IQ2 → Jetson 8GB 적재", "예정"],
]
for r in pipe:
    ws3.append(r)
    for c in range(1, 5):
        ws3.cell(row=ws3.max_row, column=c).border = BORDER
        ws3.cell(row=ws3.max_row, column=c).alignment = WRAP
widths(ws3, [20, 22, 50, 22])

# ---------------- Sheet 4: Q&A 정리 ----------------
ws4 = wb.create_sheet("Q&A")
ws4["A1"] = "질의응답 정리 (이번 세션)"
ws4["A1"].font = TITLE_FONT
ws4.merge_cells("A1:B1")
ws4.append([])
ws4.append(["질문", "핵심 답변"])
style_header(ws4, 3, 2)
qa = [
    ["Minitron 경량화 구현 방식?",
     "보정셋(한국어)으로 활성값 중요도 측정 → MLP intermediate 뉴런을 상위 70%만 남기고 텐서 실제 슬라이싱(차원 축소). fused gate_up_proj은 gate·up 같은 뉴런 동시 제거."],
    ["프루닝-디스틸-양자화 순서 맞나?",
     "논문 순서는 teacher correction → 프루닝 → 디스틸. 양자화는 논문에 없음(우리가 Jetson용으로 맨 뒤 추가). 양자화 먼저 하면 디스틸 회복이 망가짐."],
    ["teacher correction은 프루닝 전에?",
     "그렇다. 원본을 우리 데이터로 파인튜닝 → corrected 모델 생성. 이 모델이 ‘프루닝 대상(student 시작점)’ + ‘teacher’ 양쪽 역할."],
    ["teacher=corrected Phi-4면 student는?",
     "같은 corrected Phi-4를 프루닝한 작은 버전. 별도 모델 아님. 살아남은 weight를 물려받아 시작 → 디스틸 회복이 빠름."],
    ["디스틸은 student를 학습하나?",
     "그렇다. teacher는 frozen, student만 weight 업데이트. loss=KL(teacher logits ‖ student logits). soft label까지 흡수."],
    ["teacher를 비-Phi 계열로 쓰면 문제?",
     "토크나이저/vocab 불일치 때문. logit KD는 같은 어휘 분포 위에서 KL 계산 → vocab 다르면 정렬 불가. 입력 토큰 분절도 달라짐. 그래서 같은 계열 필수."],
    ["논문에 양자화 언급 있나?",
     "방법론엔 없음. 본문 5.3절 Fig.10에서 압축모델을 H100 FP8로 ‘추론 속도 벤치’한 게 유일. 압축 기법으로는 안 씀."],
    ["프루닝 테스트 진행 방식?",
     "한국어 256×1024 보정셋 → 3종 산출물(activation/magnitude width, depth) 생성 → eval_compare.py로 KMMLU 일괄 비교."],
    ["Block Influence가 뭐야?",
     "레이어가 hidden state를 얼마나 바꾸나. BI=1-cos(입력,출력). 입력≈출력(cos≈1)이면 BI≈0=잉여→제거. Phi-4는 레이어0 중요, 27~37 잉여."],
    ["모델 평가 방식?",
     "lm-eval로 multiple_choice. 프롬프트+A/B/C/D 4개 logprob 계산→argmax→정답 비교. 과목별 acc를 표본 가중평균. 0-shot."],
    ["lm-evaluation-harness 쓴 거 맞나?",
     "맞다. EleutherAI lm-eval 0.4.12. 우리는 simple_evaluate 얇은 래퍼만 작성. 데이터·채점·집계는 전부 lm-eval."],
    ["tasks=['kmmlu']면 자동 로딩?",
     "그렇다. task 정의의 dataset_path: HAERAE-HUB/KMMLU 따라 HF에서 자동 다운로드·캐시·평가. 데이터 경로 직접 줄 필요 없음."],
    ["20개만 했다는 의미?",
     "limit=20 = 과목당 20문항. 45과목×20=900문항만 평가(전체 ~35,000의 2~3%). 빠른 상대비교용, 표본 작아 절대값엔 노이즈."],
    ["eval_kmmlu 호출부 위치?",
     "운영 호출은 scripts/eval_compare.py:30 한 곳(run_kmmlu). 나머지는 tests/common/test_eval_kmmlu.py(monkeypatch 테스트)."],
]
for r in qa:
    ws4.append(r)
    for c in range(1, 3):
        ws4.cell(row=ws4.max_row, column=c).border = BORDER
        ws4.cell(row=ws4.max_row, column=c).alignment = WRAP
widths(ws4, [38, 90])

out = "docs/results/2026-06-17-pruning-results.xlsx"
wb.save(out)
print("saved:", out)
